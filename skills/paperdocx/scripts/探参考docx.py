# -*- coding: utf-8 -*-
r"""★★ 从一份**参考 docx** 里把排版参数读出来 —— 这套参数的来源就是它。

    py -3.12 -X utf8 探参考docx.py 参考稿.docx              # 全读
    py -3.12 -X utf8 探参考docx.py 参考稿.docx --只 表格
    py -3.12 -X utf8 探参考docx.py 参考稿.docx --配置       # 直接吐出 排版配置.py 片段

**只读，绝不改参考文件。**

## 为什么这一步最重要

`排版配置.py` 里的每个数字 —— 页边距 2.30 cm、正文 11 pt、行距 1.5、
表头上线 sz=12、公式编号列 1.80 cm —— **都是从真实投稿稿里读出来的，不是编的**。

所以换一本刊物的标准动作是：

```
① 找一份那本刊物的、排版好的稿子（最好是自己或同事投过的）
② 跑这个脚本，把参数读出来
③ 抄进 排版配置.py
④ 出一版，渲染成图，和参考稿并排看
```

**不许凭印象填参数。** 「RSE 大概是 12 pt 吧」这种猜法，出来的稿子一眼就不对。

## 实战：2026-10-16 读了两份，各取一半

| 参考稿 | 学到了什么 |
|---|---|
| `V34_RSE英文稿_排版版.docx` | 页面、页边距、字号、标题层级、行距、段前段后、三线表的线宽、公式双列表的结构 |
| `Manuscript_ISPRS.docx` | ★ **表格每个单元格都居中** —— 作者说「V34 里面图表的排版不好」就是这一条 |

V34 的表只有**表头**居中，数据行是左/右对齐，读起来散；
ISPRS 的表**整表每格 `jc=center`**，紧凑得多。两份各取所长才是最后的样子。

★ 这件事说明：**参考稿不必只看一份，也不必全盘照抄。**
一份看整体版式，一份看局部（图表、公式），把好的那一半挑出来。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def 净(el) -> str:
    """把 XML 打印出来，去掉一堆 xmlns 和 w: 前缀，好读。"""
    if el is None:
        return "（无）"
    s = etree.tostring(el, pretty_print=True).decode()
    s = re.sub(r'\sxmlns(:\w+)?="[^"]*"', "", s)
    return re.sub(r"\b(w|w14|w15|mc|m|r):", "", s)


def 磅(x):
    return None if x is None else round(x.pt, 2)


# ══════════════════════════════════════════════ 页面
def 读页面(doc):
    print("━━━ 页面 ━━━")
    for i, s in enumerate(doc.sections):
        print(f"[节 {i}] 纸张 {s.page_width.cm:.2f} × {s.page_height.cm:.2f} cm")
        print(f"   页边距  上 {s.top_margin.cm:.2f}  下 {s.bottom_margin.cm:.2f}  "
              f"左 {s.left_margin.cm:.2f}  右 {s.right_margin.cm:.2f} cm")
        print(f"   ★ 正文宽 = {s.page_width.cm - s.left_margin.cm - s.right_margin.cm:.2f} cm"
              f"   ← 表宽、图宽、公式列宽都按它算")
        xml = s._sectPr.xml
        for 名, 标 in (("lnNumType", "行号"), ("cols", "分栏")):
            m = re.search(r"<w:%s[^>]*/?>" % 名, xml)
            print(f"   {标}：{m.group(0) if m else '无'}")


# ══════════════════════════════════════════════ 样式
def 读样式(doc, 详细=False):
    print("\n━━━ 用到的段落样式（按出现次数） ━━━")
    用 = {}
    for p in doc.paragraphs:
        用[p.style.name] = 用.get(p.style.name, 0) + 1
    for k, v in sorted(用.items(), key=lambda x: -x[1]):
        print(f"   {k:28s} {v} 段")

    print("\n━━━ 样式定义 ━━━")
    print(f"   {'样式':26s} {'字体':20s} {'号':>5s} {'粗':>5s} {'对齐':>8s} "
          f"{'行距':>6s} {'段前':>6s} {'段后':>6s} {'首行缩进':>8s}")
    for s in doc.styles:
        if s.name not in 用 and not (s.name or "").lower().startswith(
                ("heading", "normal", "title", "caption", "body")):
            continue
        try:
            f, pf = s.font, getattr(s, "paragraph_format", None)
            行 = [f"   {s.name[:26]:26s}", f"{str(f.name)[:20]:20s}",
                 f"{str(磅(f.size)):>5s}", f"{str(f.bold):>5s}"]
            if pf is not None:
                行 += [f"{str(pf.alignment):>8s}", f"{str(pf.line_spacing):>6s}",
                      f"{str(磅(pf.space_before)):>6s}", f"{str(磅(pf.space_after)):>6s}",
                      f"{(str(round(pf.first_line_indent.cm, 2)) if pf.first_line_indent else '-'):>8s}"]
            print(" ".join(行))
        except Exception as e:
            print(f"   {s.name:26s} <读不出来：{e}>")

    # ★ 主题字：最隐蔽的坑，单独报
    d = doc.styles.element.find(W + "docDefaults")
    主题 = re.findall(r'(ascii|hAnsi|eastAsia|cs)Theme="([^"]+)"',
                     etree.tostring(doc.styles.element).decode())
    if 主题:
        print(f"\n   ⚠ 样式表里有 {len(主题)} 处**主题字**（asciiTheme 等）。"
              f"\n     主题字优先级高于显式字体名 —— 抄参数时要看 theme1.xml 里主题字到底是什么，"
              f"\n     自己出稿时则要**先删掉 *Theme 属性再设实名**，否则标题会变成无衬线。")
    if 详细 and d is not None:
        print("\n━━━ docDefaults ━━━")
        print(净(d)[:1600])


# ══════════════════════════════════════════════ 正文样例
def 读段落(doc, n=40):
    print(f"\n━━━ 前 {n} 段（看版式长什么样） ━━━")
    for i, p in enumerate(doc.paragraphs[:n]):
        pf = p.paragraph_format
        r = p.runs[0] if p.runs else None
        ri = (f" [run {r.font.name}/{磅(r.font.size)}/b={r.font.bold}"
              f"/sup={r.font.superscript}]") if r else ""
        图 = " 【图】" if p._p.find(".//" + qn("a:blip")) is not None else ""
        式 = " 【公式】" if p._p.find(".//" + MATH + "oMath") is not None else ""
        print(f"{i:3d} <{p.style.name[:20]:20s}> 对齐={str(pf.alignment):12s} "
              f"行距={str(pf.line_spacing):6s} 前={str(磅(pf.space_before)):6s} "
              f"后={str(磅(pf.space_after)):6s}{ri}{图}{式}")
        if p.text.strip():
            print(f"      {p.text.strip()[:95]}")


# ══════════════════════════════════════════════ 表格
def 读表格(doc):
    # ★ 别用 `t not in 数据表` 分类 —— python-docx 的 Table 没定义 __eq__，
    #   `in` 会退化成逐个 == 比较并全部判不等，结果两边都是全集（实测 10 张报成 10+10）。
    #   按 _tbl 的身份比。
    是公式表 = lambda t: (len(t.columns) == 2
                      and t._tbl.find(".//" + MATH + "oMath") is not None)
    数据表 = [t for t in doc.tables if not 是公式表(t)]
    公式表 = [t for t in doc.tables if 是公式表(t)]
    print(f"\n━━━ 表格：共 {len(doc.tables)} 张 = 数据表 {len(数据表)} + 公式表 {len(公式表)} ━━━")

    if 数据表:
        t = 数据表[0]
        print(f"\n【数据表样例】{len(t.rows)} 行 × {len(t.columns)} 列  "
              f"style={t.style.name if t.style else None}")
        print("--- tblPr（表宽/居中/边框/格内边距/布局）---")
        print(净(t._tbl.find(W + "tblPr")))
        g = t._tbl.find(W + "tblGrid")
        if g is not None:
            宽 = [int(c.get(qn("w:w"))) / 567 for c in g]
            print(f"--- 列宽 (cm) --- {[round(x, 2) for x in 宽]}  合计 {sum(宽):.2f}")
        for 名, ri in (("表头行", 0), ("数据行", 1), ("末行", len(t.rows) - 1)):
            if ri >= len(t.rows) or ri < 0:
                continue
            row = t.rows[ri]
            print(f"\n--- {名} trPr（cantSplit / tblHeader / jc）---")
            print(净(row._tr.find(W + "trPr")))
            c = row.cells[0]
            print(f"--- {名} 单元格0 tcPr（★ 三线表的线就在 tcBorders 里）---")
            print(净(c._tc.tcPr))
            p = c.paragraphs[0]
            print(f"--- {名} 单元格0 段落 pPr ---")
            print(净(p._p.find(W + "pPr")))
            if p.runs:
                print(f"    字体 {p.runs[0].font.name} / {磅(p.runs[0].font.size)} pt "
                      f"/ 粗={p.runs[0].font.bold}")
        print("\n★ 怎么看三线表：表头行 tcBorders 的 top（粗）+ bottom（细），"
              "末行的 bottom（粗），其余 nil。")
        print("★ 怎么看「每格都居中」：数据行的 pPr 里有没有 jc=center，"
              "tcPr 里有没有 vAlign=center。")
        print("  —— 2026-10-16 实测：V34 只有表头有，ISPRS 整表都有。这就是两份参考稿的差别。")

    if 公式表:
        t = 公式表[0]
        print(f"\n【公式表样例】{len(t.rows)} 行 × {len(t.columns)} 列")
        print("--- tblPr（★ tblBorders 应该六条全 nil）---")
        print(净(t._tbl.find(W + "tblPr")))
        g = t._tbl.find(W + "tblGrid")
        if g is not None:
            宽 = [int(c.get(qn("w:w"))) / 567 for c in g]
            print(f"--- 列宽 (cm) --- 公式列 {宽[0]:.2f}，编号列 {宽[-1]:.2f}")
        for ci in range(len(t.columns)):
            c = t.rows[0].cells[ci]
            print(f"--- 单元格{ci} tcPr --- {净(c._tc.tcPr).strip()}")
            print(f"--- 单元格{ci} pPr  --- {净(c.rows[0].cells[ci].paragraphs[0]._p.find(W+'pPr')).strip() if False else 净(c.paragraphs[0]._p.find(W+'pPr')).strip()}")
            print(f"    含 oMath: {c._tc.find('.//' + MATH + 'oMath') is not None}"
                  f"   文本: {c.text.strip()[:30]!r}")


# ══════════════════════════════════════════════ 图
def 读图(doc):
    print(f"\n━━━ 图：{len(doc.inline_shapes)} 张 ━━━")
    for i, sh in enumerate(doc.inline_shapes[:8]):
        print(f"   [{i}] {sh.width.cm:.2f} × {sh.height.cm:.2f} cm")


# ══════════════════════════════════════════════ 引用
def 读引用(doc, 路径):
    import zipfile
    with zipfile.ZipFile(路径) as z:
        xml = z.read("word/document.xml").decode("utf-8", "ignore")
        部件 = [n for n in z.namelist() if not n.startswith("word/media")]
    print("\n━━━ 引用机制 ━━━")
    print(f"   ADDIN ZOTERO 出现 {xml.count('ADDIN ZOTERO')} 次；"
          f"CSL_CITATION {xml.count('CSL_CITATION')} 次；fldChar {xml.count('fldChar')} 次")
    if xml.count("ADDIN ZOTERO") == 0:
        print("   ⇒ 这份参考稿里的引用是**死文本**，不是 Zotero 活字段。"
              "\n     （很多投稿稿都是这样：作者 Refresh 完就把域取消了）")
    print(f"   有没有 docProps/custom.xml（Zotero 认领凭据）："
          f"{'有' if 'docProps/custom.xml' in 部件 else '没有'}")
    m = re.search(r"ADDIN ZOTERO_ITEM CSL_CITATION (.{0,500})", xml)
    if m:
        print(f"\n   首个引用域码片段：\n   {m.group(1)[:460]}")


# ══════════════════════════════════════════════ 吐配置片段
def 吐配置(doc):
    r"""★★ 从**真实段落**里量，不是读样式定义。

    这是实测踩出来的：V34 那份参考稿里 `Body Text` 样式的 `size` / `line_spacing`
    全是 `None`，而真实段落的 run 是 **11 pt、行距 1.5**。
    很多投稿稿（尤其 pandoc 生成再手改的）**格式落在 run 和段落上，不在样式定义里**。
    只读样式定义会得到一串 None，照抄出来的稿子完全不对。

    ⇒ 做法：按样式名把真实段落分组，每组取**出现次数最多**的那个值（众数）。
    """
    from collections import Counter
    s = doc.sections[0]
    正文宽 = s.page_width.cm - s.left_margin.cm - s.right_margin.cm

    组: dict[str, list] = {}
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        组.setdefault(p.style.name, []).append(p)

    def 众数(值们):
        好 = [v for v in 值们 if v is not None]
        return Counter(好).most_common(1)[0][0] if 好 else None

    def 量(样式名):
        ps = 组.get(样式名)
        if not ps:
            return None
        st = next((x for x in doc.styles if x.name == 样式名), None)
        继承 = lambda 取样式, 取段: 众数([取段(p) for p in ps]) or (取样式(st) if st else None)
        对齐码 = 众数([p.paragraph_format.alignment for p in ps])
        return dict(
            n=len(ps),
            号=继承(lambda x: 磅(x.font.size),
                  lambda p: 磅(p.runs[0].font.size) if p.runs else None),
            # ★ 粗体看**全段是不是都粗**，不能只看第一个 run：
            #   正文里有「**Weighted velocity loss.** 后面接正文」这种黑体小标题开头的段，
            #   只看首 run 会把整个正文样式判成粗体（实测 V34 就这样）。
            粗=bool(众数([all(r.font.bold for r in p.runs) for p in ps if p.runs])),
            对齐={0: "left", 1: "center", 2: "right", 3: "justify"}.get(
                int(对齐码) if 对齐码 is not None else -1, "left"),
            行距=继承(lambda x: x.paragraph_format.line_spacing,
                   lambda p: p.paragraph_format.line_spacing),
            前=继承(lambda x: 磅(x.paragraph_format.space_before),
                  lambda p: 磅(p.paragraph_format.space_before)) or 0,
            后=继承(lambda x: 磅(x.paragraph_format.space_after),
                  lambda p: 磅(p.paragraph_format.space_after)) or 0,
            字体=众数([p.runs[0].font.name for p in ps if p.runs]),
        )

    def 行(键, 样式名):
        v = 量(样式名)
        if v is None:
            有 = "、".join(sorted(组)[:6])
            return (f"    # {键}：参考稿里没有「{样式名}」这个样式。"
                    f"实际用到的是：{有}")
        return (f'    "{键}": dict(号={v["号"]}, 粗={v["粗"]}, 对齐="{v["对齐"]}", '
                f'行距={v["行距"]}, 前={v["前"]}, 后={v["后"]}),'
                f'   # {样式名}，{v["n"]} 段，字体 {v["字体"]}')

    print("\n" + "=" * 78)
    print("# ↓↓↓ 抄进 排版配置.py（★ 抄完要出一版渲染成图，和参考稿并排看）")
    print("=" * 78)
    print(f'纸张 = dict(宽cm={s.page_width.cm:.2f}, 高cm={s.page_height.cm:.2f})')
    print(f'页边距 = dict(上={s.top_margin.cm:.2f}, 下={s.bottom_margin.cm:.2f}, '
          f'左={s.left_margin.cm:.2f}, 右={s.right_margin.cm:.2f})')
    print(f'正文宽cm = {正文宽:.2f}')
    print(f'加行号 = {bool(re.search("lnNumType", s._sectPr.xml))}')
    print("\n样式 = {   # ★ 从真实段落量的（众数），不是读样式定义")
    候选 = {
        "题名": ["Title"],
        "H2": ["Heading 2", "Heading 1", "标题 2"],
        "H3": ["Heading 3", "标题 3"],
        "H4": ["Heading 4", "标题 4"],
        "正文": ["Body Text", "Normal", "First Paragraph", "正文"],
        "图注": ["Caption", "Figure Caption", "题注"],
    }
    for 键, 名们 in 候选.items():
        名 = next((n for n in 名们 if n in 组), 名们[0])
        print(行(键, 名))
    print("}")
    print("\n# 参考稿里实际用到的全部样式（上面没覆盖到的自己补）：")
    for k, v in sorted(组.items(), key=lambda x: -len(x[1])):
        u = 量(k)
        print(f"#   {k:24s} {len(v):4d} 段  {u['号']} pt / 行距 {u['行距']} / "
              f"{u['对齐']} / 前{u['前']} 后{u['后']}")
    print("\n# 表格的线宽要自己从上面的 tcBorders 里抄 sz 值；")
    print("# 公式两列的宽度从上面的「公式表样例 列宽」里抄。")


# ══════════════════════════════════════════════ 主
def 主():
    参 = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not 参:
        raise SystemExit(__doc__.split("\n\n")[1].strip() + "\n\n⛔ 要给一个 docx 路径")
    路径 = Path(参[0])
    if not 路径.exists():
        raise SystemExit(f"⛔ 找不到 {路径}")
    只 = sys.argv[sys.argv.index("--只") + 1] if "--只" in sys.argv else None

    print("=" * 78)
    print(f"参考稿：{路径}  ({路径.stat().st_size / 1e6:.1f} MB)")
    print("=" * 78)
    doc = Document(str(路径))

    if "--配置" in sys.argv:
        读页面(doc)
        吐配置(doc)
        return
    跑 = lambda 名: (只 is None) or (只 == 名)
    if 跑("页面"):
        读页面(doc)
    if 跑("样式"):
        读样式(doc, 详细="--详细" in sys.argv)
    if 跑("段落"):
        读段落(doc)
    if 跑("表格"):
        读表格(doc)
    if 跑("图"):
        读图(doc)
    if 跑("引用"):
        读引用(doc, 路径)


if __name__ == "__main__":
    主()
