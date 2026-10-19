# -*- coding: utf-8 -*-
r"""英文版 md → 排版好的 Word。**全量重建**，每次重跑都从 md 出发。

    py -3.12 -X utf8 <scripts>/出Word.py                       # 出最新一版
    py -3.12 -X utf8 <scripts>/出Word.py --标红对照 上一版_英文版.md  # 改动标红
    py -3.12 -X utf8 <scripts>/出Word.py --图dpi 300 --出 _预览.docx  # 快速预览版

产物：`<稿名>_排版版.docx`

## 它做了什么

    md ──pandoc──> 骨架 docx（行内公式/粗斜体/表/图都已经是 Word 原生对象）
        ──python-docx──> 套样式、三线表、公式双列表、图缩放、
                         Zotero 活引用 + 文献表、每块埋书签锚点

★ **样式全部走 Word 样式（style），不做 run 级硬格式。**
  作者在 Word 里改一次样式就全文生效；而且 `三方同步.py` 只碰文字不碰样式，
  重建才会覆盖样式 —— 所以**作者改过样式之后就别再重建**，走同步。

## 关键约定

* 每个 md 块在 docx 里起始处埋一个书签 `PDX0001`、`PDX0002`……
  `三方同步.py` 靠它把 Word 段落认回 md 块（前缀见 排版配置.锚点前缀）。**别手动删书签。**
* 行内公式和引用在 Word 里是 OMML / Zotero 字段，**只认 md 为准**。
  在 Word 里改公式或引用，同步会报错不写。
* 参考文献走 Zotero 活字段；**不出 DOI** 的做法是把 CSL 的 `access` 宏掏空，
  **Zotero 库里的 DOI 一个都不用删**（见 references/zotero接法.md）。

## 踩过的坑

1. pandoc 默认开 `fancy_lists`，`(1) *Skeleton extraction* ...` 这种开头
   会被当成有序列表，段落变成 List 样式、前面多出编号。**必须 `-fancy_lists`。**
2. pandoc 默认开 `smart`，`Leonardo's` 的直引号会变成弯引号 `’`，
   docx 文字就和 md 不逐字相同，三方同步每次都误报「这段变了」。**必须 `-smart`。**
3. pandoc 会给 `$$...$$` 后面多塞一个空段落。重建时统一清掉「无文字无图无公式」的空段。
4. `w:pPr` / `w:tblPr` 里子元素**顺序是 schema 规定死的**，顺序错了 Word 报
   「文件已损坏」。最后统一排一遍序（见 `排子元素序`）。
"""
from __future__ import annotations

import copy
import json
import re
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

这里 = Path(__file__).resolve().parent
sys.path.insert(0, str(这里))
import 块模型 as M                                    # noqa: E402
import 排版配置 as C                                   # noqa: E402
import _zotero字段 as Z                                # noqa: E402
import docx读写 as D                                   # noqa: E402

import _项目                                           # noqa: E402
根 = _项目.找根()   # ★ 不传参数：先从当前工作目录找，再从脚本位置找
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
DRAW = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

对齐表 = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


# ══════════════════════════════════════════════════ 引用
def 找引用数据() -> Path:
    """实现在 `_项目.找引用数据`（只许有一份，别在各脚本里各写一遍）。"""
    return _项目.找引用数据(根, 这里)


class 引用库:
    def __init__(self):
        d = 找引用数据()
        self.条目 = {r["id"]: r for r in json.loads((d / "references.json").read_text("utf-8"))}
        self.索引 = json.loads((d / "zot_match_index.json").read_text("utf-8"))
        self.渲染 = {}
        # ★ 和 references.json 一样走 找引用数据()，别再单独写一串候选路径 ——
        #   脚本装进插件缓存之后，`这里/引用数据` 根本不是项目的那一份。
        映 = next((p for p in (d / "引用映射.tsv",
                             根 / "引用映射.tsv",
                             根 / "_工作区/引用映射.tsv") if p.exists()), None)
        if 映 is None:
            raise SystemExit(f"⛔ 找不到 引用映射.tsv（refN ↔ 渲染引用）。"
                             f"找过 {d}、项目根、_工作区/。"
                             f"跑 `配文献.py --写` 会生成它。")
        for line in 映.read_text("utf-8").splitlines()[1:]:
            ref, 文 = line.split("\t")
            self.渲染[ref] = 文
        self.反查 = {v: k for k, v in self.渲染.items()}
        if len(self.反查) != len(self.渲染):
            raise SystemExit("⛔ 引用映射.tsv 里有重复的渲染形式，反查会歧义")
        # 长的排前面，免得 "Liu et al., 2022" 把 "Liu, C. et al., 2023" 切一半
        形 = sorted(self.渲染.values(), key=len, reverse=True)
        单 = "|".join(re.escape(s) for s in 形)
        self.正则 = re.compile(r"\((%s)(?:; (?:%s))*\)" % (单, 单))

    def 反查并正则(self):
        """给 `块模型.md转同步串` 用的 {渲染文本: refN} + `__正则__`。"""
        d = dict(self.反查)
        d["__正则__"] = self.正则
        return d

    def 打标记(self, 文本: str) -> str:
        """正文里的渲染引用 → `[[cite:ref1;ref2]]` 标记，给 pandoc 之后替换成字段。"""
        def 换(m):
            部分 = [p.strip() for p in m.group(0)[1:-1].split(";")]
            return "[[cite:%s]]" % ";".join(self.反查[p] for p in 部分)
        return self.正则.sub(换, 文本)


# ══════════════════════════════════════════════════ 1. 造 pandoc 输入
def 造pandoc输入(块们, 库: 引用库):
    r"""返回 (md 文本, 期望结构, 公式号列表, 被引 refN 集合)

    期望结构是 [(块, 'p'|'tbl', 条数)]，后面拿它把 docx 的段落认回 md 块。
    """
    段, 期望, 公式号, 被引 = [], [], [], set()
    整句粗表题 = []
    for b in 块们:
        t = b.类型
        if t == "authors":
            正 = b.文本[len("**Authors:**"):].strip()
            段.append(正)
            期望.append((b, "p", 1))
        elif t == "affil":
            行 = [l.strip() for l in b.文本.split("\n")[1:] if l.strip()]
            段.append("\n\n".join(行))
            期望.append((b, "p", len(行)))
        elif t == "corresp":
            段.append(b.文本[len("**Correspondence:**"):].strip())
            期望.append((b, "p", 1))
        elif t == "tabcap":
            # 表题两种写法都收，**原样过**：
            #     **Table 1.** 正文        只有标签加粗（推荐，和图注一致）
            #     **Table 1. 正文**        整句加粗（V39/V40 那一版稿子就是这么写的）
            # ★ 不在这里把整句加粗改写成只粗标签。改写了 docx 的加粗结构就和 md 不一样，
            #   `三方同步.py` 每次都会把这些块报成「作者在 Word 里改过」—— 一个表题一条，
            #   7 张表就是 7 条假警报。要统一风格就去改 md，不在出 Word 时偷偷动。
            # 实测见过三种（同一个项目里就有两种）：
            #     **Table 1.** 正文                 只粗标签（推荐，和图注一致）
            #     **Table 1. 正文**                 整句粗（V39/V40 的稿子）
            #     **Table A2. 一句话.** 后面还有正文   粗到某处为止
            # 判据只要两条：开头是 `**Table N.`，后面有配对的 `**`。
            # 编号同 块模型.图表号：1 / A2 / S1 都算（补充材料是 S 开头）
            if not M.图表号.match(b.文本):
                raise SystemExit(
                    f"⛔ 表题要以 `**Table N.` 开头（块 {b.序}）：{b.文本[:80]}")
            if b.文本.count("**") < 2:
                raise SystemExit(
                    f"⛔ 表题的 `**` 没闭合（块 {b.序}）：{b.文本[:80]}")
            if not re.match(r"^\*\*(?:Table) [A-Za-z]?\d+[a-z]?\.\*\* ", b.文本):
                整句粗表题.append(b.序)
            段.append(库.打标记(b.文本))
            期望.append((b, "p", 1))
        elif t == "equation":
            体, 号, _锚 = M.拆公式块(b.文本)
            if 号 is None:
                raise SystemExit(f"⛔ 公式块没有编号（块 {b.序}）")
            公式号.append(号)
            段.append("$$\n%s\n$$" % 体)
            期望.append((b, "p", 1))
        elif t == "table":
            段.append(库.打标记(b.文本))
            期望.append((b, "tbl", 1))
        elif t == "bibph":
            段.append("ZOTEROBIBPLACEHOLDER")
            期望.append((b, "p", 1))
        else:
            段.append(库.打标记(b.文本))
            期望.append((b, "p", 1))
    文 = "\n\n".join(段) + "\n"
    for m in re.finditer(r"\[\[cite:([^\]]+)\]\]", 文):
        被引.update(m.group(1).split(";"))
    if 整句粗表题:
        print(f"  ⚠ {len(整句粗表题)} 个表题不是「只粗标签」的写法（块 {整句粗表题}），"
              "docx 里照 md 原样加粗。\n"
              "    想统一成只粗标签，把 md 改成 `**Table N.** 正文`（图注就是这么写的）；"
              "这里不替你改 —— 改了 docx 的加粗结构就和 md 不一致，三方同步会天天误报。")
    return 文, 期望, 公式号, 被引


# ══════════════════════════════════════════════════ 1.5 图：给 docx 用的降采样副本
def 备图(临: Path, 块们, dpi=None):
    r"""准备送进 docx 的图。

    ★ 默认 `排版配置.图['送入dpi'] = 0`，**用原图，一个像素都不动**
      —— 作者 2026-10-16 明确要求「图片不要压缩」。
    只有显式传 `--图dpi N` 时才在临时目录里做降采样副本（出快速预览版用），
    原图任何时候都不动。
    """
    dpi = dpi if dpi is not None else C.图.get("送入dpi")
    用到 = sorted({m.group(1) for b in 块们 if b.类型 == "image"
                 for m in [re.match(r"!\[\]\(([^)]+)\)", b.文本.strip())] if m})
    缺 = [f for f in 用到 if not (根 / f).exists()]
    if 缺:
        raise SystemExit(f"⛔ 这些图文件不存在：{缺}")
    if not dpi:
        return 根, 0, 0
    from PIL import Image
    原合计 = 新合计 = 0
    for 相对 in 用到:
        源 = 根 / 相对
        目 = 临 / 相对
        目.parent.mkdir(parents=True, exist_ok=True)
        原合计 += 源.stat().st_size
        im = Image.open(源)
        目标宽 = int(C.图["最大宽cm"] / 2.54 * dpi)
        if im.width > 目标宽:
            im = im.resize((目标宽, max(1, round(im.height * 目标宽 / im.width))),
                           Image.LANCZOS)
        im.save(目, optimize=True)
        新合计 += 目.stat().st_size
    return 临, 原合计, 新合计


# ══════════════════════════════════════════════════ 2. 页面与样式
def 设页面(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(C.纸张["宽cm"]), Cm(C.纸张["高cm"])
    s.top_margin, s.bottom_margin = Cm(C.页边距["上"]), Cm(C.页边距["下"])
    s.left_margin, s.right_margin = Cm(C.页边距["左"]), Cm(C.页边距["右"])
    # 行号：作者明确说不要。pandoc 不会加，这里显式关掉以防参考样式带进来。
    for ln in list(s._sectPr.findall(qn("w:lnNumType"))):
        s._sectPr.remove(ln)
    if C.页脚页码:
        p = s.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        Z.append_field(p, " PAGE ", "1")


def _样式(doc, 名, 基=None):
    """拿到样式；没有就新建一个段落样式。"""
    for st in doc.styles:
        if st.name == 名:
            return st
    st = doc.styles.add_style(名, 1)          # 1 = WD_STYLE_TYPE.PARAGRAPH
    if 基:
        st.base_style = doc.styles[基]
    return st


def _配(st, 规, 字体):
    f = st.font
    f.name = 字体
    f.size = Pt(规["号"])
    f.bold = 规["粗"]
    f.color.rgb = RGBColor(0, 0, 0)
    f.italic = 规.get("斜", False)
    # eastAsia 也设一遍，否则 Word 给非 ASCII 字符（× − ⁻² µ）挑中文字体，一眼看出来
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    # ★★ 必须先删掉 *Theme 属性。pandoc 的 Heading / Title 样式带
    #    `asciiTheme="majorHAnsi"`，**主题字优先于显式字体名**，只设 w:ascii
    #    是没用的 —— 实测现象是正文 Times New Roman、标题却是一水儿的无衬线。
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), 字体)

    pf = st.paragraph_format
    pf.alignment = 对齐表[规["对齐"]]
    pf.line_spacing = 规["行距"]
    pf.space_before = Pt(规["前"])
    pf.space_after = Pt(规["后"])
    pf.keep_with_next = bool(规.get("随后"))
    pf.first_line_indent = Cm(0)
    if 规.get("悬挂cm"):
        pf.left_indent = Cm(规["悬挂cm"])
        pf.first_line_indent = Cm(-规["悬挂cm"])


def _按id(doc, style_id):
    """★ 不要用 `doc.styles['Heading 2']` 按名字取 —— pandoc 生成的 styles.xml 里
    这些样式的 `w:name` 不一定是 python-docx 认的内置名，会 KeyError。按 styleId 取。"""
    for st in doc.styles:
        if st.style_id == style_id:
            return st
    return None


def 建样式(doc):
    """建/改全部段落样式，返回 {md 块类型: 样式对象}。

    样式全走 Word 样式，不做 run 级硬格式 —— 作者改一次样式就全文生效。
    """
    字, 规 = C.西文字体, C.样式
    # docDefaults 和其它样式里也带主题字，一并换成实名，否则「没被我覆盖到的地方」
    # 还是走主题的无衬线字
    主题换实名 = {"w:asciiTheme": "w:ascii", "w:hAnsiTheme": "w:hAnsi",
              "w:eastAsiaTheme": "w:eastAsia", "w:cstheme": "w:cs"}
    for rf in doc.styles.element.iter(qn("w:rFonts")):
        for 主, 实 in 主题换实名.items():
            if rf.get(qn(主)) is not None:
                del rf.attrib[qn(主)]
                rf.set(qn(实), 字)
    _配(doc.styles["Normal"], 规["正文"], 字)          # Normal 是根，先压成正文

    def 造(名, 规键, 基="Normal"):
        st = _样式(doc, 名, 基)
        _配(st, 规[规键] if isinstance(规键, str) else 规键, 字)
        return st

    表 = {}
    表["h1"] = 造("Title", "题名", None)
    表["authors"] = 造("Authors", "作者")
    表["affil"] = 造("Affiliation", "单位")
    表["corresp"] = 造("Correspondence", "通讯")
    for n, k in ((2, "H2"), (3, "H3"), (4, "H4")):
        st = _按id(doc, "Heading%d" % n)
        if st is None:
            raise SystemExit("⛔ pandoc 的 docx 里没有 Heading%d 样式" % n)
        _配(st, 规[k], 字)
        表["h%d" % n] = st
    正 = 造("Body Text", "正文")
    造("First Paragraph", "正文")                      # pandoc 给标题后第一段用这个
    表["para"] = 表["equation"] = 表["bibph"] = 正
    表["keywords"] = 造("Keywords", "关键词")
    表["image"] = 造("Figure", "图")
    表["figcap"] = 造("Figure Caption", "图注")
    表["tabcap"] = 造("Table Caption", "表题")
    表["tabnote"] = 造("Table Note", "表注")
    造("Bibliography", "文献")

    表规 = dict(号=C.表格["字号"], 粗=False,
              对齐="center" if C.表格["水平居中"] else "left",
              行距=C.表格["行距"], 前=C.表格["格内段前pt"], 后=C.表格["格内段后pt"])
    造("Compact", 表规)                                 # pandoc 表格单元格默认用它
    造("Table Cell", 表规)
    造("Table Header", dict(表规, 粗=True))
    造("Equation Number", dict(号=规["正文"]["号"], 粗=False, 对齐="center", 行距=1.0,
                             前=C.公式["段前pt"], 后=C.公式["段后pt"]))
    return 表


# ══════════════════════════════════════════════════ 3. 三线表
def _边(tc, **边框):
    pr = tc.get_or_add_tcPr()
    old = pr.find(qn("w:tcBorders"))
    if old is not None:
        pr.remove(old)
    el = OxmlElement("w:tcBorders")
    for 名 in ("top", "left", "bottom", "right"):
        b = OxmlElement("w:" + 名)
        粗 = 边框.get(名)
        if 粗:
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), str(粗))
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "000000")
        else:
            b.set(qn("w:val"), "nil")
        el.append(b)
    pr.append(el)


def _垂直居中(tc):
    pr = tc.get_or_add_tcPr()
    for old in list(pr.findall(qn("w:vAlign"))):
        pr.remove(old)
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), "center")
    pr.append(v)


# Times New Roman 各类字符的宽度，单位 em。按真实字宽分档，别一律按 0.5：
# ★ 实测教训：按 0.5 em 估的时候，表头 "LAI" / "NDVI" 被 Word 拆成 "LA↵I" / "ND↵VI"
#   —— 大写字母在 TNR 里是 0.66~0.72 em，按 0.5 估出来的列宽根本不够。
def 文宽em(s: str) -> float:
    w = 0.0
    for ch in s:
        if ch in "ilj.,;:'’!|()[]":
            w += 0.30
        elif ch in " ":
            w += 0.25
        elif ch.isdigit():
            w += 0.50
        elif ch.isupper():
            w += 0.70
        elif ch in "mwMW—–":
            w += 0.90
        elif ch.islower():
            w += 0.48
        else:
            w += 0.58
    return w


def 算列宽(格, 字号):
    r"""按内容估一套列宽；放不下返回 None（外层换小一档字号再试）。"""
    边 = C.表格["左右边距twips"] / 1440 * 2.54          # twips → cm
    em = 字号 / 28.35 * 1.06                           # pt → cm，再留 6% 余量
    列 = list(zip(*格))
    下限 = [max((max((文宽em(w) for w in str(c).split()), default=1) for c in col),
             default=1) * em + 2 * 边 for col in 列]
    最长串 = [max((文宽em(str(c)) for c in col), default=1) for col in 列]
    意愿 = [max(d, min(s, 20.0) * em + 2 * 边) for d, s in zip(下限, 最长串)]
    可用 = C.正文宽cm
    if sum(下限) > 可用:
        return None
    # 把「意愿宽」缩放到正好铺满可用宽，但谁都不许低于下限；不够的从宽的里面匀
    宽 = list(意愿)
    for _ in range(40):
        差 = 可用 - sum(宽)
        if abs(差) < 0.01:
            break
        可动 = [i for i in range(len(宽)) if 差 > 0 or 宽[i] - 下限[i] > 0.01]
        if not 可动:
            break
        池 = sum(宽[i] - 下限[i] for i in 可动) if 差 < 0 else sum(宽[i] for i in 可动)
        if 池 <= 0:
            break
        for i in 可动:
            份 = (宽[i] - 下限[i]) if 差 < 0 else 宽[i]
            宽[i] = max(下限[i], 宽[i] + 差 * 份 / 池)
    return 宽


def 配数据表(doc, tbl, 块序):
    """套三线表：全格水平+垂直居中、表头加粗、宽度按内容分。

    放不下时**先收格内左右边距，再缩字号** —— 顺序反过来的代价实测过：
    12 张表里有 5 张被缩到 7~8.5 pt；换成先收边距，11 张都留在 9 pt，
    只有 11 列的表 4 降到 8 pt。
    """
    表 = next(t for t in doc.tables if t._tbl is tbl)
    格 = [[c.text.strip() for c in row.cells] for row in 表.rows]
    原边距 = C.表格["左右边距twips"]
    号, 边距, 宽 = None, 原边距, None
    for 试号 in C.表格["缩字号档"]:
        for 试边 in (原边距, 60, 40):
            C.表格["左右边距twips"] = 试边
            宽 = 算列宽(格, 试号)
            if 宽:
                号, 边距 = 试号, 试边
                break
        if 宽:
            break
    C.表格["左右边距twips"] = 原边距
    if not 宽:
        号, 边距 = C.表格["缩字号档"][-1], 40
        print(f"   ⚠ 块 {块序} 的表在最小字号 {号}pt + 最小边距下仍放不下，"
              f"交给 Word 自动换行")
        宽 = [C.正文宽cm / len(格[0])] * len(格[0])

    pr = tbl.find(qn("w:tblPr"))
    if pr is None:
        pr = OxmlElement("w:tblPr")
        tbl.insert(0, pr)
    for 名 in ("w:tblStyle", "w:tblW", "w:jc", "w:tblBorders", "w:tblCellMar", "w:tblLayout"):
        for old in list(pr.findall(qn(名))):
            pr.remove(old)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(int(sum(宽) * 567)))
    w.set(qn("w:type"), "dxa")
    pr.append(w)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    pr.append(jc)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")      # 不固定的话 Word 会无视我算的列宽
    pr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for 名, v in (("top", C.表格["格边距twips"]), ("left", 边距),
                 ("bottom", C.表格["格边距twips"]), ("right", 边距)):
        e = OxmlElement("w:" + 名)
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    pr.append(mar)

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for x in 宽:
        g = OxmlElement("w:gridCol")
        g.set(qn("w:w"), str(int(x * 567)))
        grid.append(g)
    tbl.insert(list(tbl).index(pr) + 1, grid)

    末 = len(表.rows) - 1
    # 短表整张跟着表题走，别让 Word 在表题底下只放两行就翻页（实测表 4 就这样）。
    整张不拆 = len(表.rows) <= C.表格["整张不拆行数上限"]
    for ri, row in enumerate(表.rows):
        trPr = row._tr.get_or_add_trPr()
        for old in (list(trPr.findall(qn("w:cantSplit")))
                    + list(trPr.findall(qn("w:tblHeader")))):
            trPr.remove(old)
        trPr.append(OxmlElement("w:cantSplit"))
        if ri == 0 and C.表格["表头重复"]:
            trPr.append(OxmlElement("w:tblHeader"))
        for ci, cell in enumerate(row.cells):
            _边(cell._tc,
                top=C.表格["上线粗"] if ri == 0 else None,
                bottom=(C.表格["表头下线粗"] if ri == 0 else
                        C.表格["下线粗"] if ri == 末 else None))
            if C.表格["垂直居中"]:
                _垂直居中(cell._tc)
            tcw = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                cell._tc.tcPr.append(tcw)
            tcw.set(qn("w:w"), str(int(宽[min(ci, len(宽) - 1)] * 567)))
            tcw.set(qn("w:type"), "dxa")
            for p in cell.paragraphs:
                p.style = "Table Header" if (ri == 0 and C.表格["表头加粗"]) else "Table Cell"
                p.paragraph_format.keep_with_next = 整张不拆 and ri < 末
                # ★ pandoc 按 md 的 `--:` / `:-:` 在单元格段落上写了**直接的 w:jc**，
                #   它压过样式里的居中。删掉它，让「Table Cell」样式说了算 ——
                #   这样作者以后想改对齐，改一次样式就全表生效。
                ppr = p._p.find(qn("w:pPr"))
                if ppr is not None:
                    for old in list(ppr.findall(qn("w:jc"))):
                        ppr.remove(old)
            # 表里的行内公式（$R_b$ 这种）跟着表的字号走。OMML 的 m:r 不吃段落样式，
            # 不显式设 sz 的话它按正文 11 pt 画，在 8~9 pt 的表里一眼就大一圈。
            for mr in cell._tc.iter(qn("m:r")):
                rpr = mr.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr")
                    mr.insert(0, rpr)
                for 名 in ("w:sz", "w:szCs"):
                    for old in list(rpr.findall(qn(名))):
                        rpr.remove(old)
                    e = OxmlElement(名)
                    e.set(qn("w:val"), str(int(号 * 2)))
                    rpr.append(e)
            for p in cell.paragraphs:
                if 号 != C.表格["字号"]:
                    for r in p.runs:
                        r.font.size = Pt(号)       # 这一张表缩了字号，只能落到 run 上
    return 号


# ══════════════════════════════════════════════════ 4. 公式 → 双列隐藏边框表
def 配公式(doc, p, 号):
    """把一个含 oMathPara 的段落换成 1×2 隐藏边框表：左公式、右编号，都居中。"""
    表 = doc.add_table(rows=1, cols=2)
    p.addprevious(表._tbl)

    pr = 表._tbl.find(qn("w:tblPr"))
    for 名 in ("w:tblStyle", "w:tblW", "w:jc", "w:tblBorders", "w:tblCellMar", "w:tblLayout"):
        for old in list(pr.findall(qn(名))):
            pr.remove(old)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(int(C.正文宽cm * 567)))
    w.set(qn("w:type"), "dxa")
    pr.append(w)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    pr.append(jc)
    bd = OxmlElement("w:tblBorders")
    for 名 in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + 名)
        e.set(qn("w:val"), "nil")
        bd.append(e)
    pr.append(bd)
    mar = OxmlElement("w:tblCellMar")
    for 名 in ("left", "right"):
        e = OxmlElement("w:" + 名)
        e.set(qn("w:w"), "0")
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    pr.append(mar)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    pr.append(lay)
    grid = 表._tbl.find(qn("w:tblGrid"))
    表._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for x in (C.公式["公式列cm"], C.公式["编号列cm"]):
        g = OxmlElement("w:gridCol")
        g.set(qn("w:w"), str(int(x * 567)))
        grid.append(g)
    表._tbl.insert(list(表._tbl).index(pr) + 1, grid)

    左, 右 = 表.cell(0, 0), 表.cell(0, 1)
    左._tc.remove(左.paragraphs[0]._p)
    左._tc.append(copy.deepcopy(p))                      # 整段（含 oMathPara）搬过去
    for lp in 左.paragraphs:
        lp.style = "Equation Number"
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for mp in 左._tc.iter(qn("m:oMathPara")):            # OMML 自己也要居中
        mpr = mp.find(qn("m:oMathParaPr"))
        if mpr is None:
            mpr = OxmlElement("m:oMathParaPr")
            mp.insert(0, mpr)
        j = mpr.find(qn("m:jc"))
        if j is None:
            j = OxmlElement("m:jc")
            mpr.append(j)
        j.set(qn("m:val"), "center")
    右.paragraphs[0].style = "Equation Number"
    右.paragraphs[0].text = "(%s)" % 号
    右.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for c, x in ((左, C.公式["公式列cm"]), (右, C.公式["编号列cm"])):
        _垂直居中(c._tc)
        _边(c._tc)
        tcw = c._tc.tcPr.find(qn("w:tcW"))
        if tcw is None:
            tcw = OxmlElement("w:tcW")
            c._tc.tcPr.append(tcw)
        tcw.set(qn("w:w"), str(int(x * 567)))
        tcw.set(qn("w:type"), "dxa")
    p.getparent().remove(p)
    return 表._tbl


# ══════════════════════════════════════════════════ 5. 走一遍正文，贴样式 + 埋锚点
def 有东西(p) -> bool:
    return bool("".join(p.itertext()).strip()) or \
        p.find(".//" + DRAW + "blip") is not None or \
        p.find(".//" + MATH + "oMath") is not None


def 埋锚(p_or_tbl, 名, 书签号):
    s = OxmlElement("w:bookmarkStart")
    s.set(qn("w:id"), str(书签号))
    s.set(qn("w:name"), 名)
    e = OxmlElement("w:bookmarkEnd")
    e.set(qn("w:id"), str(书签号))
    if p_or_tbl.tag == qn("w:tbl"):
        p_or_tbl.addprevious(s)
        p_or_tbl.addnext(e)
    else:
        ppr = p_or_tbl.find(qn("w:pPr"))
        p_or_tbl.insert(1 if ppr is not None else 0, s)
        p_or_tbl.append(e)


def 配正文(doc, 期望, 公式号, 样式表):
    body = doc.element.body
    # pandoc 会在 $$ 后面留空段，先清掉；图片段看着是空的，别误杀
    for p in list(body.findall(qn("w:p"))):
        if not 有东西(p):
            body.remove(p)

    节点 = [x for x in body if x.tag in (qn("w:p"), qn("w:tbl"))]
    i, 号迭 = 0, iter(公式号)
    表字号, 上一块是表 = {}, False
    for b, 要, 条 in 期望:
        if i >= len(节点):
            raise SystemExit(f"⛔ docx 段落不够用了：块 {b.序}（{b.类型}）没有对应节点")
        首 = 节点[i]
        if 要 == "tbl":
            if 首.tag != qn("w:tbl"):
                raise SystemExit(f"⛔ 块 {b.序} 应该是表，docx 里却是段落："
                                 f"{''.join(首.itertext())[:70]}")
            表字号[b.序] = 配数据表(doc, 首, b.序)
            埋锚(首, b.锚, b.序)
            i += 1
            上一块是表 = True
            continue
        if b.类型 == "equation":
            if 首.find(".//" + MATH + "oMathPara") is None:
                raise SystemExit(f"⛔ 块 {b.序} 应该是独立公式，docx 里不是："
                                 f"{''.join(首.itertext())[:70]}")
            新 = 配公式(doc, 首, next(号迭))
            埋锚(新, b.锚, b.序)
            i += 1
            上一块是表 = False
            continue
        for k in range(条):
            节 = 节点[i + k]
            if 节.tag != qn("w:p"):
                raise SystemExit(f"⛔ 块 {b.序} 要 {条} 个段落，第 {k} 个却是表")
            P = next(p for p in doc.paragraphs if p._p is 节)
            P.style = 样式表[b.类型]
            if 上一块是表 and k == 0:
                # 表的下边框和后面那段之间留口气（Word 的表没有「段后」，
                # 只能加在后面这一段的段前上）
                P.paragraph_format.space_before = Pt(
                    C.样式[{"tabnote": "表注"}.get(b.类型, "正文")]["前"]
                    + C.表格["表后空pt"])
            if b.类型 == "image":
                for sh in doc.inline_shapes:
                    if sh._inline.getparent().getparent().getparent() is 节:
                        比 = min(Cm(C.图["最大宽cm"]) / sh.width,
                                Cm(C.图["最大高cm"]) / sh.height, 1)
                        sh.width, sh.height = int(sh.width * 比), int(sh.height * 比)
        埋锚(节点[i], b.锚, b.序)
        i += 条
        上一块是表 = False
    # 节点用光了才算对上；剩下的说明 pandoc 多吐了东西，不许装没看见
    剩 = [x for x in 节点[i:]]
    if 剩:
        raise SystemExit(f"⛔ docx 里多出 {len(剩)} 个节点没有对应的 md 块，"
                         f"第一个：{''.join(剩[0].itertext())[:70]}")
    return 表字号


# ══════════════════════════════════════════════════ 6. Zotero 活引用 + 文献表
def 插引用(doc, 库: 引用库):
    n = 0
    for p in Z.all_paragraphs(doc):
        for run in list(p.runs):
            匹 = list(re.finditer(r"\[\[cite:([^\]]+)\]\]", run.text))
            if not 匹:
                continue
            锚, 原 = run._r, run.text
            run.text = ""
            rpr = Z.rpr_like(锚)
            游 = 0
            for m in 匹:
                if m.start() > 游:
                    锚 = _插文本(锚, 原[游:m.start()], rpr)
                键 = m.group(1).split(";")
                # ★ 显示文本直接用 md 里的渲染形式（`引用映射.tsv`），不用
                #   Z.format_inline 现编。原因：四组重名作者（Liu C./Liu X. 等）
                #   靠 CSL 的 disambiguate 消歧成 "Liu, X. et al., 2023"，
                #   现编只会得到 "Liu et al., 2023"，没 Refresh 之前看着就是错的。
                显 = "(" + "; ".join(库.渲染[k] for k in 键) + ")"
                项 = [{"id": 库.索引[k]["itemID"], "uris": [库.索引[k]["uri"]],
                      "itemData": 库.条目[k]} for k in 键]
                # "Yang and Liu (Yang and Liu, 2026)" → Word 里显示 "Yang and Liu (2026)"
                # ★ 作者名要从**渲染形式**里切（去掉结尾的 ", 年"），不能用
                #   Z.author_string 现编：消歧过的形式是 "Liu, C. et al., 2023"，
                #   现编只给 "Liu et al."，`endswith` 对不上，suppress-author
                #   就不生效，Word 里印成 "Liu, C. et al. (Liu, C. et al., 2023)"。
                if len(键) == 1:
                    形 = 库.渲染[键[0]]
                    作名 = re.sub(r",\s*(?:19|20)\d{2}[a-z]?$", "", 形)
                    if 原[: m.start()].rstrip().endswith(作名) or \
                            原[: m.start()].rstrip().endswith(
                                Z.author_string(库.条目[键[0]])):
                        项[0]["suppress-author"] = True
                        显 = "(%s)" % Z.year_of(库.条目[键[0]])
                锚 = Z.insert_field_after(锚, Z.citation_instr(项, 显), 显, rpr)
                游 = m.end()
                n += 1
            if 游 < len(原):
                _插文本(锚, 原[游:], rpr)
    return n


def _插文本(锚, 文, rpr):
    r = OxmlElement("w:r")
    r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = 文
    r.append(t)
    锚.addnext(r)
    return r


def 做无DOI样式(装进Zotero=None):
    r"""把 elsevier-harvard.csl 复制一份、把 access 宏掏空 → 不出 DOI / URL。
    **Zotero 库里的 DOI 字段一个都没动**，只是这个样式不去印它。

    还会把样式装进 Zotero 的 `styles/` 目录（只新增一个文件，不碰库）。
    不装的话：docx 里预渲染的文献表是没 DOI 的，但作者一在 Word 里点 Refresh，
    Zotero 会按它自己装着的 elsevier-harvard 重排，DOI 又回来了。
    ★ 装完要**重启 Zotero** 才认得这个样式。
    """
    d = 找引用数据()
    # ★ 底版 CSL 的文件名**不写死** —— 换刊物就是换这个文件。
    #   配置里指定了就用配置的；没指定就取 引用数据/ 里唯一那个不带 -nodoi 的 .csl。
    名 = C.文献.get("底版CSL")
    源路 = (d / 名) if 名 else None
    if 源路 is None or not 源路.exists():
        候选 = [p for p in sorted(d.glob("*.csl")) if "-nodoi" not in p.stem]
        if not 候选:
            raise SystemExit(
                f"⛔ {d} 里没有 .csl 引用样式。\n"
                "   去 Zotero 数据目录的 styles/ 里拷一份刊物指定的样式进来，\n"
                "   或在 排版配置.文献 里写 底版CSL=\"xxx.csl\"。")
        源路 = 候选[0]
        if len(候选) > 1:
            print(f"  ⚠ 引用数据/ 里有 {len(候选)} 个 .csl，用了 {源路.name}。"
                  f"要指定就在 排版配置.文献 里写 底版CSL")
    出 = d / (源路.stem + "-nodoi.csl")
    源 = 源路.read_text("utf-8")
    空宏 = '<macro name="access"><text value=""/></macro>'
    文, n = re.subn(r'<macro name="access">.*?</macro>', 空宏, 源, flags=re.S)
    # ★ 刊名缩写：elsevier-harvard 对期刊论文写的是
    #   `<text variable="container-title" form="short"/>`，citeproc 就去取条目的
    #   「期刊缩写」字段，印出 "Remote Sens. Environ."。作者稿子里写的是全称，
    #   **28 条有出入里有一半是这一条造成的**。把 form="short" 去掉就全称了。
    #   （只改这份派生样式，Zotero 库里的「期刊缩写」字段一个都不动。）
    if C.文献.get("刊名全称"):
        文, m = re.subn(r'(<text variable="container-title")\s+form="short"',
                        r"\1", 文)
        if m == 0:
            print("  ⚠ 这个样式里没有 container-title form=\"short\"，"
                  "刊名全称这一项没起作用")
    if n == 0 and re.search(r'\b(DOI|URL)\b', 源):
        # 有的样式不走 access 宏，直接在 bibliography 里印 DOI
        raise SystemExit(
            f"⛔ {源路.name} 里没有 `access` 宏，但正文里出现了 DOI/URL。\n"
            "   这个样式印 DOI 的方式不一样，去看它的 <bibliography> 怎么写的，\n"
            "   手改一份去掉 DOI 的存成 <名>-nodoi.csl，脚本会直接用。")
    # id 和 rel="self" 的 href 必须一致，不一致 Zotero 装样式时会拒
    # id 和 rel="self" 的 href 必须一致，不一致 Zotero 装样式时会拒
    文 = re.sub(r"<id>[^<]*</id>", "<id>%s</id>" % C.文献["样式ID"], 文, count=1)
    文 = re.sub(r'<link href="[^"]*" rel="self"/>',
                '<link href="%s" rel="self"/>' % C.文献["样式ID"], 文, count=1)
    文 = re.sub(r"<title>[^<]*</title>", "<title>%s</title>" % C.文献["样式名"], 文, count=1)
    文 = re.sub(r"<title-short>[^<]*</title-short>", "", 文, count=1)
    出.write_text(文, encoding="utf-8")

    if 装进Zotero:
        # ★ 数据目录统一走 `_项目.Zotero数据目录()`（配置 > 环境变量 > ~/Zotero）。
        #   这儿原来自己写了一套，只看环境变量和 ~/Zotero，结果把样式装进了
        #   一个 2022 年废弃的 profile，不报错、也不生效。
        数据 = _项目.Zotero数据目录()
        # ★ 文件名跟着源样式走，**不许写死**。
        #   写死的代价实测过：跑 demo 时把真实项目那份 elsevier-harvard-nodoi.csl
        #   覆盖成了 demo 的内容。
        目标 = 数据 / "styles" / 出.name
        if 数据.exists():
            目标.parent.mkdir(exist_ok=True)
            旧 = 目标.read_text("utf-8") if 目标.exists() else None
            if 旧 != 文:
                目标.write_text(文, encoding="utf-8")
                print(f"  ↳ 样式已装进 Zotero：{目标}（**要重启 Zotero 才认**）")
        else:
            print(f"  ⚠ 找不到 Zotero 数据目录 {数据}，样式没装进去；"
                  f"作者在 Word 里 Refresh 会用回带 DOI 的 elsevier-harvard")
    return 出


def 插文献(doc, 库: 引用库, 被引, csl):
    # ★ 「正文引到的」不等于「文献表要印的」。只在补充材料里引的条目，作者照样
    #   放在正文的 References 里 —— 走 排版配置.文献["额外收录"]，不然会被静默丢掉。
    额外 = set(C.文献.get("额外收录") or ())
    缺 = sorted(额外 - set(库.条目))
    if 缺:
        raise SystemExit(f"⛔ 额外收录里的 {缺} 不在 references.json 里")
    要 = (set(被引) | 额外) - set(C.文献["排除"])
    收 = [库.条目[k] for k in 库.条目 if k in 要]
    丢 = sorted(被引 & set(C.文献["排除"]))
    if 丢:
        raise SystemExit(f"⛔ {丢} 被排除，正文里却还引着它。先把正文的引用去掉。")
    # references.json 里有 HTML 实体（`Photogrammetric Engineering &amp; Remote Sensing`），
    # 不脱一层的话参考文献里就印出 "&amp;" 四个字母。实测见过。
    def 脱实体(x):
        if isinstance(x, str):
            return (x.replace("&amp;", "&").replace("&lt;", "<")
                    .replace("&gt;", ">").replace("&quot;", '"').replace("&#38;", "&"))
        if isinstance(x, list):
            return [脱实体(i) for i in x]
        if isinstance(x, dict):
            return {k: 脱实体(v) for k, v in x.items()}
        return x

    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        (d / "bib.md").write_text('---\nnocite: "@*"\n---\n', encoding="utf-8")
        (d / "refs.json").write_text(json.dumps(脱实体(收), ensure_ascii=False),
                                     encoding="utf-8")
        subprocess.run(["pandoc", str(d / "bib.md"), "--citeproc",
                        "--bibliography=" + str(d / "refs.json"),
                        "--csl=" + str(csl), "-o", str(d / "bib.docx")], check=True)
        bib = Document(d / "bib.docx")
        条 = [copy.deepcopy(p._p) for p in bib.paragraphs if p.text.strip()]

    # citeproc 会把题名做成蓝色超链接。投稿稿里参考文献要是普通黑字，
    # 把 w:hyperlink 拆掉只留里面的 run，顺手去掉 Hyperlink 字符样式和颜色。
    for bp in 条:
        for link in list(bp.iter(qn("w:hyperlink"))):
            父 = link.getparent()
            i = list(父).index(link)
            for 子 in reversed(list(link)):
                父.insert(i, 子)
            父.remove(link)
        for rpr in bp.iter(qn("w:rPr")):
            for 名 in ("w:color", "w:u", "w:rStyle"):
                for old in list(rpr.findall(qn(名))):
                    rpr.remove(old)

    放好 = False
    for p in list(doc.paragraphs):
        if "ZOTEROBIBPLACEHOLDER" not in p.text:
            continue
        锚 = p._p
        for bp in 条:
            ppr = bp.find(qn("w:pPr"))
            if ppr is None:
                ppr = OxmlElement("w:pPr")
                bp.insert(0, ppr)
            for old in list(ppr.findall(qn("w:pStyle"))):
                ppr.remove(old)
            ps = OxmlElement("w:pStyle")
            ps.set(qn("w:val"), "Bibliography")
            ppr.insert(0, ps)
            锚.addnext(bp)
            锚 = bp
        # ★ Zotero 书目字段跨段落：begin/instr/separate 塞进**第一条**的开头、
        #   end 挂在**最后一条**的末尾。早先的写法是另起两个空段落放 begin 和 end，
        #   结果 "References" 标题和第一条之间空出一大截。
        首 = 条[0]
        首ppr = 首.find(qn("w:pPr"))
        插点 = list(首).index(首ppr) + 1 if 首ppr is not None else 0
        for r in reversed(Z.make_field_runs(Z.bibliography_instr(), "")[:3]):
            首.insert(插点, r)
        条[-1].append(Z.make_field_runs("", "")[-1])
        p._p.getparent().remove(p._p)
        放好 = True
    if not 放好:
        raise SystemExit("⛔ 没找到文献占位符 ZOTEROBIBPLACEHOLDER")
    return len(条)


def 装Zotero偏好(路径):
    """Word 自定义文档属性里的 Zotero 会话信息。没有它，Zotero 不认这是它的文档。"""
    pref = ('<?xml version="1.0"?><data data-version="3" zotero-version="10.0">'
            '<session id="%s"/>'
            '<style id="%s" hasBibliography="1" bibliographyStyleHasBeenSet="1"/>'
            '<prefs><pref name="fieldType" value="Field"/>'
            '<pref name="automaticJournalAbbreviations" value="false"/></prefs></data>'
            % (C.文献.get("会话ID", "PaperDocx"), C.文献["样式ID"]))
    ns = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
    vt = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
    custom = etree.Element("{%s}Properties" % ns, nsmap={None: ns, "vt": vt})
    prop = etree.SubElement(custom, "{%s}property" % ns,
                            fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}", pid="2",
                            name="ZOTERO_PREF_1")
    etree.SubElement(prop, "{%s}lpwstr" % vt).text = pref
    with zipfile.ZipFile(路径) as z:
        文件 = {n: z.read(n) for n in z.namelist()}
    文件["docProps/custom.xml"] = etree.tostring(custom, xml_declaration=True,
                                                encoding="UTF-8", standalone=True)
    ct = etree.fromstring(文件["[Content_Types].xml"])
    if not any(x.get("PartName") == "/docProps/custom.xml" for x in ct):
        etree.SubElement(ct, "{http://schemas.openxmlformats.org/package/2006/"
                             "content-types}Override", PartName="/docProps/custom.xml",
                         ContentType="application/vnd.openxmlformats-officedocument."
                                     "custom-properties+xml")
        文件["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True,
                                                    encoding="UTF-8", standalone=True)
    rels = etree.fromstring(文件["_rels/.rels"])
    if not any(x.get("Target") == "docProps/custom.xml" for x in rels):
        etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/"
                               "relationships}Relationship", Id="rIdZoteroPreferences",
                         Type="http://schemas.openxmlformats.org/officeDocument/2006/"
                              "relationships/custom-properties",
                         Target="docProps/custom.xml")
        文件["_rels/.rels"] = etree.tostring(rels, xml_declaration=True,
                                            encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(路径, "w", zipfile.ZIP_DEFLATED) as z:
        for n, v in 文件.items():
            z.writestr(n, v)


# ══════════════════════════════════════════════════ 6.5 和上一版 diff，改过的地方标红
def 标红(doc, 块们, 上一版: Path, 库: 引用库):
    r"""把本版相对 `上一版` 改过 / 新增的文字在 docx 里标成深红。

        py -3.12 -X utf8 出Word.py --标红对照 C_英文稿V41_英文版.md

    做法：两版的块按**规约后的文本**用 difflib 对齐（中间插删过段落也对得上），
    改过的块再做**按词的 diff**，只把新增 / 替换的词标红，没动的词保持黑色。

    ★ 公式、引用、图片不标红：它们在 docx 里是 OMML / Zotero 字段 / 图，
      `写段` 只能原样搬过去（见 docx读写 的说明）。整块新增时才整块标红。
    """
    import difflib
    旧块 = M.读英文版(上一版)
    条目id表 = {v["itemKey"]: k for k, v in 库.索引.items()}
    书签 = D.书签表(doc)

    可标 = {"h1", "h2", "h3", "h4", "para", "keywords", "figcap", "tabcap",
           "tabnote", "authors", "corresp", "table"}
    归 = lambda b: M.压空白(M.规约(b.文本))
    sm = difflib.SequenceMatcher(None, [归(b) for b in 旧块], [归(b) for b in 块们])
    配对 = {}                      # 新块序 -> 旧块文本（None 表示整块是新的）
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        for k in range(j1, j2):
            # replace 段里按位置配一配，insert 段没有对应的旧块
            偏 = k - j1
            配对[块们[k].序] = 旧块[i1 + 偏].文本 if (tag == "replace" and i1 + 偏 < i2) else None

    def 划红(旧串, 新串):
        """按词 diff；新增 / 替换的部分包进 ⟪r⟫…⟪/r⟫。哨兵本身不参与比对。"""
        切 = lambda s: [x for x in re.split(r"(\s+|⟪/?[bisr]⟫|⟦m\d+⟧|⟦c:[^⟧]*⟧)", s) if x]
        a, b = 切(旧串), 切(新串)
        出, 红 = [], False
        for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b).get_opcodes():
            要红 = tag in ("replace", "insert")
            for t in b[j1:j2]:
                if M.哨兵.fullmatch(t):          # 粗体/公式/引用标记原样放，不染色
                    if 红:
                        出.append(M.红关)
                        红 = False
                    出.append(t)
                    continue
                if 要红 and not 红 and t.strip():
                    出.append(M.红开)
                    红 = True
                elif not 要红 and 红:
                    出.append(M.红关)
                    红 = False
                出.append(t)
        if 红:
            出.append(M.红关)
        return "".join(出)

    改块 = 全新 = 0
    for b in 块们:
        if b.类型 not in 可标 or b.序 not in 配对:
            continue
        el = 书签.get(b.锚)
        if el is None:
            continue
        旧文 = 配对[b.序]
        if b.类型 == "table":
            旧格 = M.解析表(旧文)[1] if 旧文 else None
            新格 = M.解析表(b.文本)[1]
            trs = el.findall(D.W + "tr")
            for ri, 行 in enumerate(新格):
                if ri >= len(trs):
                    break
                tcs = trs[ri].findall(D.W + "tc")
                for ci, 格 in enumerate(行):
                    if ci >= len(tcs):
                        break
                    旧格文 = (旧格[ri][ci] if 旧格 and ri < len(旧格)
                            and ci < len(旧格[ri]) else "")
                    新 = M.md转同步串(格, 库.反查并正则())[0]
                    if M.压空白(新) == M.压空白(M.md转同步串(旧格文, 库.反查并正则())[0]):
                        continue
                    ps = tcs[ci].findall(D.W + "p")
                    if ps:
                        D.写段(ps[0], M.红开 + 新 + M.红关, 条目id表)
            改块 += 1
            continue
        前缀, 正 = "", b.文本
        m = re.match(r"^(#{1,4} )", 正)
        if m:
            前缀, 正 = m.group(1), 正[m.end():]
        else:
            for p in ("**Authors:**", "**Correspondence:**"):
                if 正.startswith(p):
                    前缀, 正 = p + " ", 正[len(p):].lstrip()
        新串 = M.md转同步串(正, 库.反查并正则())[0]
        if 旧文 is None:
            D.写段(el, M.红开 + 新串 + M.红关, 条目id表)
            全新 += 1
            continue
        旧正 = 旧文[len(前缀):] if 前缀 and 旧文.startswith(前缀) else 旧文
        旧串 = M.md转同步串(旧正, 库.反查并正则())[0]
        D.写段(el, 划红(旧串, 新串), 条目id表)
        改块 += 1
    return 改块, 全新


# ══════════════════════════════════════════════════ 7. 子元素排序（不排 Word 会说文件损坏）
顺序 = {
    "w:pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr "
             "suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap "
             "overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
             "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc "
             "textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr "
             "sectPr pPrChange",
    "w:rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss "
             "imprint noProof snapToGrid vanish webHidden color spacing w kern position sz "
             "szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout "
             "specVanish oMath",
    "w:tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize "
               "tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook "
               "tblCaption tblDescription tblPrChange",
    "w:tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection "
              "tcFitText vAlign hideMark",
    "w:trPr": "cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader "
              "tblCellSpacing jc hidden",
    "w:tcBorders": "top left start bottom right end insideH insideV tl2br tr2bl",
    "w:tblBorders": "top left start bottom right end insideH insideV",
}


def 排子元素序(doc):
    for tag, 名们 in 顺序.items():
        前缀 = tag.split(":")[0]
        秩 = {qn(前缀 + ":" + n): i for i, n in enumerate(名们.split())}
        for el in doc.element.iter(qn(tag)):
            el[:] = sorted(el, key=lambda c: 秩.get(c.tag, 999))


# ══════════════════════════════════════════════════ 主
def 取参(名, 缺省=None):
    return sys.argv[sys.argv.index(名) + 1] if 名 in sys.argv else 缺省


def 主():
    # --源 / --出 可以给绝对路径（做小样排查的时候用）
    _根, _版, _英, _对, _排 = _项目.三件套(根)          # 永远出最新那一版
    源 = Path(取参("--源", _英))
    if not 源.is_absolute():
        源 = 根 / 源
    目标 = Path(取参("--出", _排.name if 源 == _英 else
                     源.name.replace(".md", ".docx")))
    if not 目标.is_absolute():
        目标 = 根 / 目标
    if not 源.exists():
        raise SystemExit(f"⛔ 找不到 {源}")

    库 = 引用库()
    块们 = M.读英文版(源)
    文, 期望, 公式号, 被引 = 造pandoc输入(块们, 库)

    临 = Path(tempfile.mkdtemp(prefix="v41word_"))
    图dpi = 取参("--图dpi")
    图根, 图原, 图新 = 备图(临, 块们, int(图dpi) if 图dpi else None)
    (临 / "in.md").write_text(文, encoding="utf-8")
    subprocess.run(
        ["pandoc", str(临 / "in.md"),
         # -fancy_lists：别把 "(1) ..." 当有序列表；-smart：别把直引号变弯，
         # 否则 docx 文字和 md 不逐字相同，三方同步每次误报
         "--from=markdown+tex_math_dollars-fancy_lists-smart",
         "--to=docx", "--resource-path=" + str(图根), "-o", str(临 / "base.docx")],
        check=True)

    doc = Document(临 / "base.docx")
    设页面(doc)
    样式表 = 建样式(doc)
    表字号 = 配正文(doc, 期望, 公式号, 样式表)
    引用数 = 插引用(doc, 库)
    # ★ 没有文献占位块就不插文献表 —— **补充材料**就是这种：正文照样有引用
    #   （要变成 Zotero 活字段），但参考文献表在正文稿里，自己这边没有。
    #   以前无条件插，补充材料直接报「没找到文献占位符 ZOTEROBIBPLACEHOLDER」出不来。
    有文献表 = any(b.类型 == "bibph" for b in 块们)
    # ★ 以前这里写死 `做无DOI样式()`，排版配置.文献["装进Zotero"] 根本没人读 ——
    #   配了 True 也不会装进 Zotero，作者一 Refresh DOI 全回来。
    文献数 = (插文献(doc, 库, 被引, 做无DOI样式(C.文献.get("装进Zotero")))
            if 有文献表 else 0)
    红 = None
    对照 = 取参("--标红对照")
    if 对照:
        p = Path(对照)
        红 = 标红(doc, 块们, p if p.is_absolute() else 根 / p, 库)
    排子元素序(doc)
    doc.save(目标)
    装Zotero偏好(目标)

    缩 = {s: n for s, n in 表字号.items() if n != C.表格["字号"]}
    print(f"✓ {目标.name}  {目标.stat().st_size / 1e6:.1f} MB")
    if 图原:
        print(f"  ⚠ 这一版的图降采样到 {图dpi} dpi（**原图未动**）："
              f"{图原/1e6:.0f} MB → {图新/1e6:.0f} MB —— 这是预览版，别当交付稿")
    else:
        print("  图用原图，未压缩")
    print(f"  md 块 {len(块们)}：{dict(Counter(b.类型 for b in 块们))}")
    print(f"  活引用字段 {引用数} 处 / 文献 {文献数} 条（被引 {len(被引)}，"
          f"排除 {','.join(C.文献['排除'])}）"
          + ("" if 有文献表 else "  ← 这份稿子没有文献占位块，没插文献表"))
    print(f"  公式 {len(公式号)} / 表 {len(表字号)} / 图 {sum(1 for b in 块们 if b.类型=='image')}")
    if 缩:
        print(f"  为放下而缩了字号的表（块序:字号）：{缩}")
    if 红:
        print(f"  ★ 已和 {Path(对照).name} diff 标红：改过 {红[0]} 块、整块新增 {红[1]} 块")
    print(f"  锚点 {块们[0].锚}…{块们[-1].锚}；行号 {'开' if C.加行号 else '关'}")


if __name__ == "__main__":
    主()
